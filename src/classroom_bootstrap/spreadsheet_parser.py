from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable



COURSE_SHEET_NAME = "Courses"
COMPONENT_SHEET_NAME = "Components"
DEFAULT_COURSE_STATE = "PROVISIONED"
ALLOWED_COURSE_STATES = {"PROVISIONED", "ACTIVE", "ARCHIVED", "DECLINED", "SUSPENDED"}
ALLOWED_COMPONENT_TYPES = {"TEACHER", "STUDENT", "ALIAS", "TOPIC", "ANNOUNCEMENT", "COURSEWORK", "MATERIAL"}


@dataclass(frozen=True)
class ValidationIssue:
    sheet: str
    row_number: int
    message: str
    level: str = "error"


@dataclass(frozen=True)
class CourseRecord:
    row_number: int
    name: str
    section: str | None
    description: str | None
    owner_email: str
    course_state: str
    classroom_id: str | None = None
    classroom_link: str | None = None


@dataclass(frozen=True)
class ComponentRecord:
    row_number: int
    course_name: str
    type: str
    title: str
    description: str | None


@dataclass
class ParseResult:
    courses: list[CourseRecord] = field(default_factory=list)
    components: list[ComponentRecord] = field(default_factory=list)
    issues: list[ValidationIssue] = field(default_factory=list)

    @property
    def has_errors(self) -> bool:
        return any(issue.level == "error" for issue in self.issues)

    def grouped_issues(self) -> dict[str, dict[int, list[ValidationIssue]]]:
        grouped: dict[str, dict[int, list[ValidationIssue]]] = {}
        for issue in self.issues:
            grouped.setdefault(issue.sheet, {}).setdefault(issue.row_number, []).append(issue)
        return grouped


class SpreadsheetParser:
    """Parse course/bootstrap xlsx files into typed records with grouped validation feedback."""

    course_headers = {
        "name": "name",
        "section": "section",
        "description": "description",
        "owner email": "owner_email",
        "course state": "course_state",
        "classroom id": "classroom_id",
        "classroom link": "classroom_link",
    }
    required_course_headers = {"name", "owner email", "course state"}

    component_headers = {
        "course name": "course_name",
        "type": "type",
        "title": "title",
        "description": "description",
    }
    required_component_headers = {"course name", "type", "title"}

    def parse_file(self, xlsx_path: str | Path) -> ParseResult:
        load_workbook, _ = _require_openpyxl()
        workbook = load_workbook(xlsx_path, data_only=True)
        result = ParseResult()

        self._parse_courses(workbook, result)
        self._parse_components(workbook, result)

        return result

    def _parse_courses(self, workbook, result: ParseResult) -> None:
        if COURSE_SHEET_NAME not in workbook.sheetnames:
            result.issues.append(
                ValidationIssue(
                    sheet=COURSE_SHEET_NAME,
                    row_number=1,
                    message=f"Missing required sheet '{COURSE_SHEET_NAME}'.",
                )
            )
            return

        sheet = workbook[COURSE_SHEET_NAME]
        header_map = self._header_map(sheet, COURSE_SHEET_NAME, self.course_headers, self.required_course_headers, result)
        if not header_map:
            return

        for row_number in range(2, sheet.max_row + 1):
            row = self._row_data(sheet, row_number, header_map)
            if self._is_blank_row(row.values()):
                continue

            name = normalize_value(row.get("name"))
            owner_email = normalize_value(row.get("owner_email"))
            section = normalize_value(row.get("section")) or None
            description = normalize_value(row.get("description")) or None
            course_state = normalize_value(row.get("course_state"), upper=True)
            classroom_id = normalize_value(row.get("classroom_id")) or None
            classroom_link = normalize_value(row.get("classroom_link")) or None

            if not name:
                result.issues.append(self._error(COURSE_SHEET_NAME, row_number, "Course name is required."))
            if not owner_email:
                result.issues.append(self._error(COURSE_SHEET_NAME, row_number, "Owner email is required."))

            if not course_state:
                course_state = DEFAULT_COURSE_STATE
                result.issues.append(
                    ValidationIssue(
                        sheet=COURSE_SHEET_NAME,
                        row_number=row_number,
                        level="warning",
                        message=(
                            f"Course state was blank; defaulted to '{DEFAULT_COURSE_STATE}'. "
                            "Set 'course state' explicitly if you need a different lifecycle state."
                        ),
                    )
                )
            elif course_state not in ALLOWED_COURSE_STATES:
                result.issues.append(
                    self._error(
                        COURSE_SHEET_NAME,
                        row_number,
                        f"Invalid course state '{course_state}'. Allowed states: {', '.join(sorted(ALLOWED_COURSE_STATES))}.",
                    )
                )

            if any(i.sheet == COURSE_SHEET_NAME and i.row_number == row_number and i.level == "error" for i in result.issues):
                continue

            result.courses.append(
                CourseRecord(
                    row_number=row_number,
                    name=name,
                    section=section,
                    description=description,
                    owner_email=owner_email,
                    course_state=course_state,
                    classroom_id=classroom_id,
                    classroom_link=classroom_link,
                )
            )

    def _parse_components(self, workbook, result: ParseResult) -> None:
        if COMPONENT_SHEET_NAME not in workbook.sheetnames:
            result.issues.append(
                ValidationIssue(
                    sheet=COMPONENT_SHEET_NAME,
                    row_number=1,
                    message=f"Missing required sheet '{COMPONENT_SHEET_NAME}'.",
                )
            )
            return

        sheet = workbook[COMPONENT_SHEET_NAME]
        header_map = self._header_map(
            sheet,
            COMPONENT_SHEET_NAME,
            self.component_headers,
            self.required_component_headers,
            result,
        )
        if not header_map:
            return

        for row_number in range(2, sheet.max_row + 1):
            row = self._row_data(sheet, row_number, header_map)
            if self._is_blank_row(row.values()):
                continue

            course_name = normalize_value(row.get("course_name"))
            type_ = normalize_value(row.get("type"), upper=True)
            title = normalize_value(row.get("title"))
            description = normalize_value(row.get("description")) or None

            if not course_name:
                result.issues.append(self._error(COMPONENT_SHEET_NAME, row_number, "Course name is required."))
            if not title:
                result.issues.append(self._error(COMPONENT_SHEET_NAME, row_number, "Title is required."))

            if not type_:
                type_ = "MATERIAL"
                result.issues.append(
                    ValidationIssue(
                        sheet=COMPONENT_SHEET_NAME,
                        row_number=row_number,
                        level="warning",
                        message="Type was blank; defaulted to 'MATERIAL'.",
                    )
                )
            elif type_ not in ALLOWED_COMPONENT_TYPES:
                result.issues.append(
                    self._error(
                        COMPONENT_SHEET_NAME,
                        row_number,
                        f"Invalid type '{type_}'. Allowed types: {', '.join(sorted(ALLOWED_COMPONENT_TYPES))}.",
                    )
                )

            if any(i.sheet == COMPONENT_SHEET_NAME and i.row_number == row_number and i.level == "error" for i in result.issues):
                continue

            result.components.append(
                ComponentRecord(
                    row_number=row_number,
                    course_name=course_name,
                    type=type_,
                    title=title,
                    description=description,
                )
            )

    def _header_map(self, sheet, sheet_name: str, allowed_headers: dict[str, str], required: set[str], result: ParseResult) -> dict[str, int]:
        header_map: dict[str, int] = {}
        for index, cell in enumerate(sheet[1], start=1):
            normalized = normalize_header(cell.value)
            if normalized in allowed_headers:
                header_map[allowed_headers[normalized]] = index

        normalized_required = {allowed_headers[name] for name in required}
        missing = sorted(normalized_required - set(header_map))
        if missing:
            missing_message = ", ".join(missing)
            result.issues.append(
                self._error(
                    sheet_name,
                    1,
                    f"Header row is missing required column(s): {missing_message}. Please add them to row 1.",
                )
            )
            return {}

        return header_map

    @staticmethod
    def _row_data(sheet, row_number: int, header_map: dict[str, int]) -> dict[str, str | None]:
        return {name: sheet.cell(row=row_number, column=col).value for name, col in header_map.items()}

    @staticmethod
    def _is_blank_row(values: Iterable[object]) -> bool:
        return all(not normalize_value(v) for v in values)

    @staticmethod
    def _error(sheet: str, row_number: int, message: str) -> ValidationIssue:
        return ValidationIssue(sheet=sheet, row_number=row_number, message=message, level="error")



def _require_openpyxl():
    try:
        from openpyxl import Workbook, load_workbook
    except ModuleNotFoundError as exc:
        raise ModuleNotFoundError(
            "openpyxl is required for spreadsheet parsing/export. Install it with `pip install openpyxl`."
        ) from exc
    return load_workbook, Workbook

def normalize_header(value: object) -> str:
    normalized = normalize_value(value)
    if not normalized:
        return ""
    return " ".join(normalized.replace("_", " ").split()).lower()


def normalize_value(value: object, *, upper: bool = False) -> str:
    if value is None:
        return ""
    text = " ".join(str(value).strip().split())
    return text.upper() if upper else text


def format_issues_for_instructors(issues: list[ValidationIssue]) -> str:
    if not issues:
        return "No validation issues found."

    grouped: dict[str, dict[int, list[ValidationIssue]]] = {}
    for issue in issues:
        grouped.setdefault(issue.sheet, {}).setdefault(issue.row_number, []).append(issue)

    lines: list[str] = []
    for sheet_name in sorted(grouped):
        lines.append(f"{sheet_name} sheet:")
        for row_number in sorted(grouped[sheet_name]):
            lines.append(f"  Row {row_number}:")
            for issue in grouped[sheet_name][row_number]:
                prefix = "Warning" if issue.level == "warning" else "Error"
                lines.append(f"    - {prefix}: {issue.message}")
        lines.append("")

    return "\n".join(lines).strip()


def export_updated_course_list(courses: list[CourseRecord], output_path: str | Path) -> Path:
    """Write a clean course roster xlsx that includes generated IDs/links for instructor review."""

    _, Workbook = _require_openpyxl()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = COURSE_SHEET_NAME

    headers = [
        "Name",
        "Section",
        "Description",
        "Owner Email",
        "Course State",
        "Classroom ID",
        "Classroom Link",
        "Source Row",
    ]
    sheet.append(headers)

    for course in courses:
        sheet.append(
            [
                course.name,
                course.section or "",
                course.description or "",
                course.owner_email,
                course.course_state,
                course.classroom_id or "",
                course.classroom_link or "",
                course.row_number,
            ]
        )

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output
